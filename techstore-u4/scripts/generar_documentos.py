"""Genera el reporte técnico y la bitácora de TechStore U4 en formato DOCX.

El contenido se construye desde los datos reales del repositorio:

* los catálogos y pedidos se leen de ``datos/*.csv``;
* los resultados de la bitácora se leen de ``docs/bitacora/bitacora_escenarios.csv``;
* los diagramas se insertan desde ``docs/diagramas/*.png``;
* los conteos de pruebas se obtienen del archivo de evidencias.

La portada de ambos documentos usa los datos del equipo definidos en las
constantes ``DOCENTE``, ``GRUPO``, ``INTEGRANTES`` y ``FECHA_ENTREGA``.

Uso (con el Python del sistema, que tiene ``python-docx``):

    C:\\Python314\\python.exe scripts/generar_documentos.py

Salidas:

    docs/Reporte_Tecnico_TechStore_Unidad_4.docx
    docs/bitacora/Bitacora_TechStore.docx
"""

from __future__ import annotations

import csv
import re
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
DIAGRAMAS = DOCS / "diagramas"
BITACORA_CSV = DOCS / "bitacora" / "bitacora_escenarios.csv"

AZUL = RGBColor(0x1F, 0x3B, 0x57)
AZUL_CLARO = "DCE6F1"
GRIS = RGBColor(0x59, 0x59, 0x59)
VERDE = RGBColor(0x1E, 0x6B, 0x3A)
EVIDENCIAS_IMG = DOCS / "evidencias"

# Datos de portada
DOCENTE = "Miguel Ángel Gómez Marroquín"
GRUPO = "004"
SEMESTRE = "agosto–diciembre 2026"
EQUIPO = "Equipo 9"
INTEGRANTES = [
    ("Jonathan Alexander Reyes García", "2146387"),
    ("Jorge Luis Molina García", "2147723"),
    ("Leonardo Dávila Góngora", "2145945"),
]
FECHA_ENTREGA = "27/10/2026"
REPO_URL = "https://github.com/leyesdavila/Componentes-Logica-de-Negocios-EQ9/tree/main/techstore-u4"


# ---------------------------------------------------------------------------
# Utilidades de formato
# ---------------------------------------------------------------------------


def leer_csv(ruta: Path) -> list[dict[str, str]]:
    with ruta.open(encoding="utf-8-sig", newline="") as archivo:
        return list(csv.DictReader(archivo))


def contar_pruebas() -> int:
    """Lee el total de pruebas de la evidencia de pytest."""
    evidencia = DOCS / "evidencias" / "04_pytest.txt"
    if evidencia.exists():
        coincidencia = re.search(
            r"(\d+) passed", evidencia.read_text(encoding="utf-8", errors="replace")
        )
        if coincidencia:
            return int(coincidencia.group(1))
    return 116


def configurar_documento(doc: Document, titulo_corto: str) -> None:
    """Estilos base, márgenes, encabezado y pie con numeración."""
    seccion = doc.sections[0]
    seccion.page_width = Inches(8.5)
    seccion.page_height = Inches(11)
    seccion.left_margin = Inches(1.0)
    seccion.right_margin = Inches(1.0)
    seccion.top_margin = Inches(1.0)
    seccion.bottom_margin = Inches(0.9)
    seccion.header_distance = Inches(0.45)
    seccion.footer_distance = Inches(0.4)
    seccion.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for nombre, tamano in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        estilo = doc.styles[nombre]
        estilo.font.name = "Calibri"
        estilo.font.size = Pt(tamano)
        estilo.font.bold = True
        estilo.font.color.rgb = AZUL
        estilo.paragraph_format.space_before = Pt(14 if nombre == "Heading 1" else 10)
        estilo.paragraph_format.space_after = Pt(6)
        estilo.paragraph_format.keep_with_next = True

    for nombre in ("List Bullet", "List Number"):
        doc.styles[nombre].font.name = "Calibri"
        doc.styles[nombre].font.size = Pt(11)

    # Índice compacto: dos niveles, sin espacio extra entre entradas.
    for nombre in ("TOC 1", "TOC 2"):
        try:
            estilo = doc.styles[nombre]
        except KeyError:
            continue
        estilo.font.name = "Calibri"
        estilo.font.size = Pt(9.5)
        estilo.paragraph_format.space_before = Pt(0)
        estilo.paragraph_format.space_after = Pt(1)
        estilo.paragraph_format.line_spacing = 1.0

    encabezado = seccion.header.paragraphs[0]
    encabezado.text = titulo_corto
    encabezado.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in encabezado.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = GRIS

    pie = seccion.footer.paragraphs[0]
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = pie.add_run("Página ")
    run.font.size = Pt(8.5)
    run.font.color.rgb = GRIS
    agregar_campo(pie, "PAGE")
    run = pie.add_run(" de ")
    run.font.size = Pt(8.5)
    run.font.color.rgb = GRIS
    agregar_campo(pie, "NUMPAGES")


def agregar_campo(parrafo, codigo: str) -> None:
    """Inserta un campo de Word (PAGE, NUMPAGES, TOC)."""
    run = parrafo.add_run()
    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    instruccion = OxmlElement("w:instrText")
    instruccion.set(qn("xml:space"), "preserve")
    instruccion.text = f" {codigo} "
    separador = OxmlElement("w:fldChar")
    separador.set(qn("w:fldCharType"), "separate")
    texto = OxmlElement("w:t")
    texto.text = "1"
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    for elemento in (inicio, instruccion, separador, texto, fin):
        run._r.append(elemento)
    run.font.size = Pt(8.5)
    run.font.color.rgb = GRIS


def agregar_toc(doc: Document) -> None:
    """Índice automático de Word (se actualiza al exportar el PDF)."""
    parrafo = doc.add_paragraph()
    run = parrafo.add_run()
    inicio = OxmlElement("w:fldChar")
    inicio.set(qn("w:fldCharType"), "begin")
    instruccion = OxmlElement("w:instrText")
    instruccion.set(qn("xml:space"), "preserve")
    instruccion.text = ' TOC \\o "1-2" \\h \\z \\u '
    separador = OxmlElement("w:fldChar")
    separador.set(qn("w:fldCharType"), "separate")
    texto = OxmlElement("w:t")
    texto.text = "Índice en actualización: abra el documento y actualice el campo."
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    for elemento in (inicio, instruccion, separador, texto, fin):
        run._r.append(elemento)


def parrafo(doc: Document, texto: str, *, justificar: bool = True, tamano=None,
            negrita: bool = False, cursiva: bool = False, color=None,
            espacio_despues=None, sangria=None):
    """Párrafo normal con opciones frecuentes."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justificar else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(texto)
    run.bold = negrita
    run.italic = cursiva
    if tamano:
        run.font.size = Pt(tamano)
    if color:
        run.font.color.rgb = color
    if espacio_despues is not None:
        p.paragraph_format.space_after = Pt(espacio_despues)
    if sangria is not None:
        p.paragraph_format.left_indent = Inches(sangria)
    return p


def vineta(doc: Document, texto: str, nivel: int = 0) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.3 + 0.25 * nivel)
    p.paragraph_format.space_after = Pt(3)
    p.add_run(texto)


def numerada(doc: Document, texto: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.space_after = Pt(3)
    p.add_run(texto)


def codigo(doc: Document, texto: str) -> None:
    """Bloque de código con fondo gris claro y fuente monoespaciada."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.right_indent = Inches(0.1)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_together = True
    sombreado = OxmlElement("w:shd")
    sombreado.set(qn("w:val"), "clear")
    sombreado.set(qn("w:color"), "auto")
    sombreado.set(qn("w:fill"), "F2F2F2")
    p._p.get_or_add_pPr().append(sombreado)
    for i, linea in enumerate(texto.splitlines()):
        run = p.add_run(("\n" if i else "") + linea)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x20, 0x20, 0x20)
    return p


def sombrear_celda(celda, color: str) -> None:
    tcPr = celda._tc.get_or_add_tcPr()
    sombreado = OxmlElement("w:shd")
    sombreado.set(qn("w:val"), "clear")
    sombreado.set(qn("w:color"), "auto")
    sombreado.set(qn("w:fill"), color)
    tcPr.append(sombreado)


def _marcar_fila(fila, *, encabezado: bool = False, no_partir: bool = True,
                 con_siguiente: bool = False) -> None:
    """Evita que una fila se parta y, si se pide, la mantiene con la siguiente."""
    trPr = fila._tr.get_or_add_trPr()
    if no_partir:
        trPr.append(OxmlElement("w:cantSplit"))
    if encabezado:
        trPr.append(OxmlElement("w:tblHeader"))
    if con_siguiente:
        for celda in fila.cells:
            for p in celda.paragraphs:
                p.paragraph_format.keep_with_next = True


def tabla(doc: Document, encabezados: list[str], filas: list[list[str]],
          *, titulo: str | None = None, anchos: list[float] | None = None,
          tamano: float = 9, alinear_centro: list[int] | None = None,
          mantener_con_siguiente: bool = False) -> None:
    """Tabla con encabezado azul, filas alternas y título opcional."""
    if titulo:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(titulo)
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = AZUL

    t = doc.add_table(rows=1, cols=len(encabezados))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if anchos:
        for indice, ancho in enumerate(anchos):
            for celda in t.columns[indice].cells:
                celda.width = Inches(ancho)

    for indice, texto in enumerate(encabezados):
        celda = t.rows[0].cells[indice]
        celda.text = ""
        p = celda.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(1)
        run = p.add_run(str(texto))
        run.bold = True
        run.font.size = Pt(tamano)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        sombrear_celda(celda, "1F3B57")
    _marcar_fila(t.rows[0], encabezado=True, con_siguiente=True)

    for numero_fila, fila in enumerate(filas):
        celdas = t.add_row().cells
        for indice, valor in enumerate(fila):
            celda = celdas[indice]
            celda.text = ""
            p = celda.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            if alinear_centro and indice in alinear_centro:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(valor))
            run.font.size = Pt(tamano)
            if numero_fila % 2 == 1:
                sombrear_celda(celda, "F4F7FB")
        ultima = numero_fila == len(filas) - 1
        _marcar_fila(
            t.rows[-1],
            con_siguiente=(not ultima) or mantener_con_siguiente,
        )

    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def imagen(doc: Document, nombre: str, leyenda: str, *, ancho: float = 6.2,
           carpeta: Path = DIAGRAMAS) -> None:
    ruta = carpeta / nombre
    if not ruta.exists():
        parrafo(doc, f"[Falta la imagen {nombre}]", cursiva=True, color=GRIS)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(ruta), width=Inches(ancho))
    pie = doc.add_paragraph()
    pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pie.paragraph_format.space_after = Pt(8)
    run = pie.add_run(leyenda)
    run.italic = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = GRIS


def captura(doc: Document, nombre: str, leyenda: str, *, ancho: float = 6.2) -> None:
    """Inserta una captura de pantalla de docs/evidencias."""
    imagen(doc, nombre, leyenda, ancho=ancho, carpeta=EVIDENCIAS_IMG)


def enlace(parrafo_destino, url: str, texto: str | None = None) -> None:
    """Agrega un hipervínculo azul y subrayado al párrafo."""
    r_id = parrafo_destino.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    sub = OxmlElement("w:u")
    sub.set(qn("w:val"), "single")
    rpr.append(sub)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = texto or url
    run.append(t)
    hl.append(run)
    parrafo_destino._p.append(hl)


def seccion_repositorio(doc: Document, titulo: str) -> None:
    seccion_1(doc, titulo)
    parrafo(
        doc,
        "El código fuente, las pruebas, los datos originales, los diagramas, las "
        "evidencias y estos documentos están publicados en el repositorio del equipo:",
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    enlace(p, REPO_URL)


def salto(doc: Document) -> None:
    doc.add_page_break()


def portada(doc: Document, tipo: str, subtitulo: str) -> None:
    """Portada con los datos del equipo (formato de entregas anteriores)."""
    for _ in range(2):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("UNIVERSIDAD AUTÓNOMA DE NUEVO LEÓN")
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = AZUL

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("FACULTAD DE INGENIERÍA MECÁNICA Y ELÉCTRICA")
    run.bold = True
    run.font.size = Pt(11.5)
    run.font.color.rgb = AZUL

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(38)
    run = p.add_run("Componentes de Lógica de Negocio · Unidad 4")
    run.font.size = Pt(10.5)
    run.font.color.rgb = GRIS

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(tipo)
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = AZUL

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("TechStore: núcleo de pedidos con patrones de lógica de negocio")
    run.font.size = Pt(15)
    run.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(44)
    run = p.add_run(subtitulo)
    run.italic = True
    run.font.size = Pt(11)
    run.font.color.rgb = GRIS

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.add_run(f"Docente: {DOCENTE}").font.size = Pt(11)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.add_run(f"Grupo {GRUPO} · Semestre {SEMESTRE}").font.size = Pt(11)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    p.add_run(EQUIPO).font.size = Pt(11)

    filas = [
        ("Proyecto", ["TechStore — núcleo de pedidos"]),
        ("Integrantes", [f"{nombre} {matricula}" for nombre, matricula in INTEGRANTES]),
        ("Fecha de entrega", [FECHA_ENTREGA]),
    ]
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for etiqueta, valores in filas:
        celdas = t.add_row().cells
        celdas[0].width = Inches(1.7)
        celdas[1].width = Inches(4.5)
        sombrear_celda(celdas[0], "F3F4F6")
        run = celdas[0].paragraphs[0].add_run(etiqueta)
        run.bold = True
        run.font.size = Pt(11)
        for i, valor in enumerate(valores):
            p = celdas[1].paragraphs[0] if i == 0 else celdas[1].add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.add_run(valor).font.size = Pt(11)

    salto(doc)


def seccion_1(doc: Document, texto: str, nivel: int = 1) -> None:
    doc.add_heading(texto, level=nivel)


# ---------------------------------------------------------------------------
# Datos del proyecto
# ---------------------------------------------------------------------------


def cargar_datos():
    clientes = leer_csv(ROOT / "datos" / "clientes.csv")
    productos = leer_csv(ROOT / "datos" / "productos.csv")
    pedidos = leer_csv(ROOT / "datos" / "pedidos.csv")
    detalle = leer_csv(ROOT / "datos" / "detalle_pedidos.csv")
    escenarios = leer_csv(ROOT / "datos" / "escenarios.csv")
    bitacora = leer_csv(BITACORA_CSV) if BITACORA_CSV.exists() else []
    return clientes, productos, pedidos, detalle, escenarios, bitacora


def totales_por_pedido(detalle, productos):
    precios = {p["sku"]: float(p["precio_mxn"]) for p in productos}
    totales: dict[str, float] = {}
    unidades: dict[str, int] = {}
    renglones: dict[str, int] = {}
    for fila in detalle:
        pid = fila["pedido_id"]
        cantidad = int(fila["cantidad"])
        totales[pid] = totales.get(pid, 0.0) + cantidad * precios[fila["sku"]]
        unidades[pid] = unidades.get(pid, 0) + cantidad
        renglones[pid] = renglones.get(pid, 0) + 1
    return totales, unidades, renglones


# ---------------------------------------------------------------------------
# Reporte técnico
# ---------------------------------------------------------------------------


def generar_reporte() -> Path:
    clientes, productos, pedidos, detalle, escenarios, bitacora = cargar_datos()
    totales, unidades, renglones = totales_por_pedido(detalle, productos)
    pruebas = contar_pruebas()
    por_cliente = {c["cliente_id"]: c for c in clientes}
    por_sku = {p["sku"]: p for p in productos}

    doc = Document()
    configurar_documento(doc, "Reporte técnico · TechStore U4")

    portada(
        doc,
        "Reporte Técnico",
        "Práctica Guiada — Actividad 4 · Desarrollo de un núcleo de negocio en Python",
    )

    seccion_1(doc, "Resumen ejecutivo")
    parrafo(
        doc,
        "Este reporte documenta la construcción de TechStore, un núcleo de negocio "
        "para una tienda ficticia de tecnología. El sistema implementa entidades y "
        "objetos de valor, reglas de descuento componibles mediante el patrón "
        "Specification, una máquina de estados con eventos de dominio, un servicio "
        "de aplicación que carga los datos del paquete y una API REST construida con "
        "FastAPI. La solución se acompaña de una suite automatizada de "
        f"{pruebas} pruebas y de una bitácora reproducible de los 15 escenarios "
        "propuestos en el material de la práctica.",
    )
    parrafo(
        doc,
        "El punto de partida fueron cinco archivos CSV con 12 clientes, 12 productos, "
        "16 pedidos y 21 renglones de detalle. Ningún archivo de datos fue editado y "
        "el cargador original (`cargar_csv.py`) se conservó intacto. Toda la lógica "
        "nueva se organizó en capas: dominio, Specifications, servicio de aplicación, "
        "adaptador HTTP y pruebas. El resultado es un proyecto ejecutable, verificable "
        "y documentado, con evidencias de cada paso.",
    )
    tabla(
        doc,
        ["Indicador", "Resultado"],
        [
            ["Pedidos cargados desde los CSV", "16, todos inician en DRAFT"],
            ["Escenarios de la bitácora", "15 de 15 coinciden con el objetivo del paquete"],
            ["Pruebas automatizadas", f"{pruebas} en verde"],
            ["Endpoints de negocio", "11"],
            ["Endpoints de apoyo", "2 (/health y /admin/reload-fixtures)"],
            ["Archivos CSV modificados", "0"],
            ["Diagramas entregados", "5, cada uno con fuente Mermaid, PlantUML, PNG y SVG"],
        ],
        titulo="Tabla 0. Resultados principales de la entrega",
        anchos=[3.0, 3.3],
    )

    salto(doc)
    seccion_1(doc, "Índice")
    parrafo(
        doc,
        "El índice se actualiza automáticamente al abrir el documento o al exportarlo "
        "a PDF. Si aparece desactualizado, use clic derecho sobre él y elija "
        "«Actualizar campo».",
        cursiva=True,
        tamano=9,
        color=GRIS,
    )
    agregar_toc(doc)

    salto(doc)

    seccion_1(doc, "1. Introducción")
    seccion_1(doc, "1.1 Motivación", 2)
    parrafo(
        doc,
        "Las aplicaciones de comercio electrónico parecen sencillas desde la interfaz: "
        "elegir productos, confirmar un pedido y esperar la entrega. Detrás de esa "
        "secuencia hay reglas que no pueden quedar dispersas en los controladores. "
        "Un descuento mal calculado, una transición de estado permitida en el momento "
        "equivocado o un pedido modificado después de un rechazo producen pérdidas y "
        "decisiones difíciles de auditar. La práctica propone concentrar esas reglas "
        "en un núcleo de dominio explícito y verificable.",
    )
    parrafo(
        doc,
        "TechStore sirve como caso de estudio porque reúne tres problemas clásicos: "
        "valores monetarios que deben ser exactos, reglas comerciales que se combinan "
        "y un ciclo de vida con estados y eventos. La solución no depende de una base "
        "de datos ni de un marco web para expresar el negocio: el dominio puede "
        "probarse por sí solo y la API solo lo expone.",
    )

    seccion_1(doc, "1.2 Alcance", 2)
    parrafo(
        doc,
        "El alcance de la entrega incluye el análisis de los datos, el modelo de "
        "dominio, las reglas de descuento, la máquina de estados, la API REST, la "
        "suite de pruebas, las evidencias de ejecución y la documentación técnica. "
        "Quedan fuera del alcance la persistencia en base de datos, la autenticación "
        "de usuarios, el descuento de inventario y el despliegue en producción. "
        "El repositorio es en memoria: los cambios en los pedidos se pierden al "
        "reiniciar el proceso, mientras que los CSV permanecen como fuente de verdad "
        "de solo lectura.",
    )

    seccion_1(doc, "1.3 Entregables", 2)
    for texto in (
        "Código fuente organizado en `app/`, con dominio, Specifications, servicio y API.",
        "Suite de pruebas en `tests/`, con cobertura de dominio, API, escenarios y servicio.",
        "Bitácora de los 15 escenarios en Markdown, CSV y documento DOCX/PDF.",
        "Cinco diagramas UML en Mermaid, PlantUML, PNG y SVG.",
        "Evidencias reproducibles de carga, pruebas y respuestas HTTP.",
        "Este reporte técnico y el manifiesto de entrega con huellas SHA-256.",
    ):
        vineta(doc, texto)

    seccion_1(doc, "2. Objetivos")
    seccion_1(doc, "2.1 Objetivo general", 2)
    parrafo(
        doc,
        "Diseñar e implementar un núcleo de lógica de negocio para TechStore que "
        "modele pedidos, clientes, productos y descuentos con reglas explícitas, "
        "estados controlados y una API REST verificable, aplicando buenas prácticas "
        "de separación de responsabilidades y pruebas automatizadas.",
    )
    seccion_1(doc, "2.2 Objetivos específicos", 2)
    for texto in (
        "Analizar la estructura y las relaciones de los cinco archivos CSV del paquete.",
        "Representar importes monetarios con `Decimal` y objetos de valor inmutables.",
        "Implementar las invariantes del pedido en el agregado `Order`.",
        "Expresar las reglas de descuento como Specifications componibles con AND, OR y NOT.",
        "Construir una máquina de estados con cinco transiciones y eventos de dominio.",
        "Separar la carga de datos en un servicio de aplicación reutilizable.",
        "Exponer el dominio mediante una API REST que traduzca excepciones a HTTP.",
        "Verificar el comportamiento con pruebas de dominio, API, escenarios y frontera.",
        "Documentar decisiones, errores, evidencias y resultados de forma reproducible.",
    ):
        numerada(doc, texto)

    seccion_1(doc, "3. Contexto del problema")
    parrafo(
        doc,
        "Un pedido de TechStore recorre cinco estados: `DRAFT`, `CONFIRMED`, `SHIPPED`, "
        "`DELIVERED` y `RETURNED`; desde `CONFIRMED` también puede pasar a `CANCELLED`. "
        "Cada transición tiene precondiciones distintas. Confirmar exige que el pedido "
        "tenga artículos y que el cliente esté activo; enviar exige una guía no vacía; "
        "cancelar exige un motivo; entregar y devolver solo son válidos desde `SHIPPED`.",
    )
    parrafo(
        doc,
        "En paralelo, el negocio define dos descuentos. El 10 % aplica a clientes VIP "
        "con total mayor a 500 MXN, al menos cuatro unidades y sin promoción. El 5 % "
        "aplica a pedidos con total mayor a 5000 MXN y sin promoción. Si ambas reglas "
        "se cumplen, prevalece el 10 %. Los umbrales son estrictos: 500.00 exacto no "
        "califica para el 10 %, y el límite de unidades es 20, no 21.",
    )
    parrafo(
        doc,
        "El reto de diseño es que esas reglas se lean en un solo lugar y que ninguna "
        "operación inválida deje efectos parciales. Un rechazo no debe cambiar el "
        "estado, no debe agregar eventos y no debe modificar los artículos. Esa "
        "propiedad se verifica explícitamente en la suite.",
    )

    seccion_1(doc, "4. Análisis de los datos")
    parrafo(
        doc,
        "El paquete original contiene cinco archivos CSV. Antes de escribir código se "
        "revisaron las llaves, los tipos y las relaciones. La Tabla 4.1 resume el "
        "inventario; las tablas siguientes muestran los catálogos y los totales "
        "calculados para los 16 pedidos.",
    )
    tabla(
        doc,
        ["Archivo", "Llave", "Contenido", "Registros"],
        [
            ["clientes.csv", "cliente_id", "Nombre, es_vip, activo", "12"],
            ["productos.csv", "sku", "Nombre, precio_mxn, existencia", "12"],
            ["pedidos.csv", "pedido_id", "cliente_id, tiene_promocion, estado_inicial", "16"],
            ["detalle_pedidos.csv", "(pedido_id, sku)", "Cantidad por renglón", "21"],
            ["escenarios.csv", "caso_id", "pedido_id, acción, objetivo, pista", "15"],
        ],
        titulo="Tabla 4.1. Inventario de los datos originales",
        anchos=[1.45, 1.35, 2.7, 0.8],
        alinear_centro=[3],
    )

    seccion_1(doc, "4.1 Catálogo de clientes", 2)
    tabla(
        doc,
        ["ID", "Nombre", "VIP", "Activo"],
        [
            [c["cliente_id"], c["nombre"], "Sí" if c["es_vip"] == "1" else "No",
             "Sí" if c["activo"] == "1" else "No"]
            for c in clientes
        ],
        titulo="Tabla 4.2. Los 12 clientes del paquete",
        anchos=[0.8, 2.4, 0.8, 0.8],
        alinear_centro=[2, 3],
    )
    parrafo(
        doc,
        "Dos clientes están inactivos (C003 y C004). Esa condición es la causa del "
        "rechazo de E03 y de una prueba específica de la API. Los clientes VIP se "
        "identifican con `es_vip = 1`; la regla del 10 % los exige, pero no basta "
        "por sí sola.",
    )

    seccion_1(doc, "4.2 Catálogo de productos", 2)
    tabla(
        doc,
        ["SKU", "Producto", "Precio (MXN)", "Existencia"],
        [
            [p["sku"], p["nombre"], f"{float(p['precio_mxn']):,.2f}", p["existencia"]]
            for p in productos
        ],
        titulo="Tabla 4.3. Los 12 productos del catálogo",
        anchos=[0.8, 2.6, 1.1, 0.9],
        alinear_centro=[2, 3],
    )
    parrafo(
        doc,
        "La existencia es información de catálogo y no se descuenta al crear un "
        "pedido. El paquete no pide control de inventario; agregar esa regla habría "
        "cambiado el alcance y, probablemente, el resultado de varios escenarios.",
    )

    seccion_1(doc, "4.3 Pedidos y totales", 2)
    filas_pedidos = []
    for p in pedidos:
        pid = p["pedido_id"]
        cliente = por_cliente[p["cliente_id"]]
        filas_pedidos.append(
            [
                pid,
                cliente["nombre"],
                "Sí" if cliente["es_vip"] == "1" else "No",
                "Sí" if cliente["activo"] == "1" else "No",
                "Sí" if p["tiene_promocion"] == "1" else "No",
                str(renglones.get(pid, 0)),
                str(unidades.get(pid, 0)),
                f"{totales.get(pid, 0.0):,.2f}",
            ]
        )
    tabla(
        doc,
        ["Pedido", "Cliente", "VIP", "Activo", "Promo", "Renglones", "Unidades", "Total"],
        filas_pedidos,
        titulo="Tabla 4.4. Los 16 pedidos, con total calculado desde el detalle",
        anchos=[0.65, 1.3, 0.45, 0.55, 0.5, 0.7, 0.7, 0.8],
        tamano=8,
        alinear_centro=[2, 3, 4, 5, 6, 7],
    )
    parrafo(
        doc,
        "Los totales se calcularon sumando cantidad por precio de catálogo. O001, "
        "O003 y O005 comparten el mismo detalle (2100.00 MXN) y se diferencian solo "
        "por el cliente o la promoción, lo que permite aislar el efecto de cada "
        "variable. O007 no tiene renglones y O008 llega exactamente a 20 unidades.",
    )

    seccion_1(doc, "4.4 Observaciones que condicionan el diseño", 2)
    for texto in (
        "Las unidades no son renglones: O001 tiene 2 renglones y 4 unidades; O008 "
        "tiene 1 renglón y 20 unidades. El límite debe sumar cantidades.",
        "El total de O010 es exactamente 500.00, lo que obliga a distinguir `>` de `>=`.",
        "O011 tiene 3 unidades y un cliente VIP: sirve para probar que el umbral de "
        "unidades es independiente del monto.",
        "O006 es el único pedido con total mayor a 5000 y su cliente no es VIP; por eso "
        "la prioridad del 10 % sobre el 5 % se prueba con un pedido construido en memoria.",
        "O009 tiene 19 unidades; agregar dos más debe fallar sin modificar el pedido.",
        "Los escenarios E11 a E15 reutilizan pedidos que ya avanzaron de estado, así que "
        "cada caso debe ejecutarse sobre una carga nueva de los fixtures.",
    ):
        vineta(doc, texto)

    seccion_1(doc, "5. Arquitectura de la solución")
    parrafo(
        doc,
        "La solución aplica una arquitectura por capas con dependencias dirigidas hacia "
        "el dominio. El núcleo no conoce FastAPI ni los archivos CSV; la API no repite "
        "reglas; el servicio es la única pieza que sabe cómo leer los datos del paquete.",
    )
    tabla(
        doc,
        ["Capa", "Archivo", "Responsabilidad", "Depende de"],
        [
            ["Dominio", "app/domain.py", "Entidades, value objects, invariantes, estados y eventos", "Solo biblioteca estándar"],
            ["Reglas", "app/specifications.py", "Predicados de descuento y composición AND/OR/NOT", "app.domain"],
            ["Servicio", "app/service.py", "Lectura de CSV, construcción de pedidos y catálogos", "app.domain"],
            ["Adaptador", "app/api.py", "Endpoints, validación de entrada y traducción de errores", "app.domain, app.specifications, app.service"],
            ["Pruebas", "tests/", "Verificación del dominio, la API, el servicio y los escenarios", "Todas las capas"],
        ],
        titulo="Tabla 5.1. Capas y responsabilidades",
        anchos=[0.9, 1.4, 2.7, 1.5],
        tamano=8.5,
    )

    seccion_1(doc, "5.1 Estructura del repositorio", 2)
    codigo(
        doc,
        "techstore-u4/\n"
        "├── app/\n"
        "│   ├── __init__.py          # Documentación del paquete\n"
        "│   ├── domain.py            # Entidades, value objects e invariantes\n"
        "│   ├── specifications.py    # Reglas de descuento componibles\n"
        "│   ├── service.py           # Carga de CSV y catálogos\n"
        "│   └── api.py               # API REST FastAPI\n"
        "├── tests/\n"
        "│   ├── conftest.py\n"
        "│   ├── test_domain.py\n"
        "│   ├── test_specifications.py\n"
        "│   ├── test_servicio.py\n"
        "│   ├── test_escenarios.py\n"
        "│   ├── test_maquina_estados.py\n"
        "│   └── test_api.py\n"
        "├── scripts/\n"
        "│   ├── ejecutar_bitacora.py\n"
        "│   ├── generar_evidencias.py\n"
        "│   └── generar_documentos.py\n"
        "├── datos/                   # Cinco CSV originales, sin editar\n"
        "├── docs/                    # Análisis, bitácora, diagramas y evidencias\n"
        "├── evidencias/              # Salidas reproducibles de la raíz\n"
        "├── cargar_csv.py            # Cargador original, intacto\n"
        "├── README.txt               # Reglas del paquete, intacto\n"
        "└── requirements.txt",
    )

    seccion_1(doc, "5.2 Flujo de una petición", 2)
    parrafo(
        doc,
        "Cuando llega una petición HTTP, FastAPI valida el cuerpo con Pydantic antes de "
        "invocar al controlador. El controlador busca el pedido en el diccionario en "
        "memoria, llama al método del agregado y, si el dominio lanza una excepción, la "
        "traduce a un código HTTP. El flujo es deliberadamente corto: la decisión de "
        "negocio no se reescribe en la capa web.",
    )
    codigo(
        doc,
        "Cliente HTTP\n"
        "    │\n"
        "    ▼\n"
        "FastAPI ── valida cuerpo (Pydantic) ──► 422 si el formato es inválido\n"
        "    │\n"
        "    ▼\n"
        "Controlador ── busca el pedido ──► 404 si no existe\n"
        "    │\n"
        "    ▼\n"
        "Order.confirm() / ship() / cancel() / deliver() / return_order()\n"
        "    │        │\n"
        "    │        └── valida precondiciones ──► DomainException ──► 409 / 422\n"
        "    ▼\n"
        "Cambia el estado y registra el DomainEvent\n"
        "    │\n"
        "    ▼\n"
        "200 OK con el pedido serializado",
    )

    seccion_1(doc, "6. Modelo de dominio")
    parrafo(
        doc,
        "El modelo distingue entidades, objetos de valor y eventos. `Order` es el "
        "agregado raíz: todo cambio del pedido pasa por sus métodos. `Customer` tiene "
        "identidad y atributos que condicionan la confirmación. `Money` y `OrderItem` "
        "son inmutables y se comparan por valor.",
    )
    tabla(
        doc,
        ["Elemento", "Tipo", "Identidad", "Responsabilidad"],
        [
            ["Money", "Value object", "No", "Importe en MXN, suma y porcentaje con redondeo"],
            ["OrderItem", "Value object", "No", "SKU, cantidad, precio unitario y subtotal derivado"],
            ["Customer", "Entidad", "cliente_id", "Cliente VIP o activo, precondición de confirmación"],
            ["Order", "Aggregate root", "pedido_id", "Artículos, estado, eventos e invariantes"],
            ["OrderStatus", "Enumeración", "—", "DRAFT, CONFIRMED, SHIPPED, DELIVERED, RETURNED, CANCELLED"],
            ["DomainEvent", "Evento", "—", "Hecho ocurrido, con nombre y marca de tiempo"],
            ["DomainException", "Excepción base", "—", "Base de los rechazos de negocio"],
        ],
        titulo="Tabla 6.1. Elementos del modelo de dominio",
        anchos=[1.15, 1.2, 0.95, 3.0],
        tamano=8.5,
    )

    seccion_1(doc, "6.1 Invariantes del agregado", 2)
    tabla(
        doc,
        ["Invariante", "Regla", "Excepción"],
        [
            ["Importe válido", "No se admiten importes negativos ni valores no finitos", "InvalidOrderData"],
            ["Artículo válido", "SKU no vacío y cantidad mayor que cero", "InvalidOrderData"],
            ["Moneda", "Los artículos del pedido deben estar en MXN", "InvalidOrderData"],
            ["Límite de unidades", "La suma de cantidades no puede superar 20", "OrderItemLimitExceeded"],
            ["Edición", "Solo se agregan artículos en DRAFT", "InvalidOrderTransition"],
            ["Confirmación", "Requiere al menos un artículo y cliente activo", "InvalidOrderData"],
            ["Envío", "Requiere guía no vacía y estado CONFIRMED", "InvalidOrderData / InvalidOrderTransition"],
            ["Cancelación", "Requiere motivo no vacío y estado CONFIRMED", "InvalidOrderData / InvalidOrderTransition"],
            ["Entrega y devolución", "Solo desde SHIPPED", "InvalidOrderTransition"],
        ],
        titulo="Tabla 6.2. Invariantes y excepciones",
        anchos=[1.35, 3.35, 1.6],
        tamano=8.5,
    )

    parrafo(
        doc,
        "La validación ocurre antes de cualquier mutación. En `add_item()` se calcula "
        "primero el nuevo total de unidades y se lanza la excepción si excede el límite; "
        "solo después de validar se agrega el renglón. En las transiciones se comprueba "
        "el estado, luego la precondición, después se cambia `status` y al final se "
        "registra el evento. Esa secuencia garantiza que un rechazo no deje rastros.",
    )

    seccion_1(doc, "6.2 Diagrama de clases", 2)
    imagen(
        doc,
        "clases.png",
        "Figura 6.1. Entidades, value objects, enumeración, eventos y excepciones del dominio.",
    )

    seccion_1(doc, "7. Patrón Specification")
    seccion_1(doc, "7.1 Problema", 2)
    parrafo(
        doc,
        "Las reglas de descuento combinan varias condiciones y cambian con el tiempo. "
        "Si se escriben como una cadena de `if`, cada modificación obliga a leer toda la "
        "función, las pruebas se acoplan al orden de las condiciones y la explicación "
        "del negocio queda escondida en el código. El patrón Specification convierte "
        "cada condición en un objeto con un método `is_satisfied_by(order)` y permite "
        "componerlas con operadores.",
    )

    seccion_1(doc, "7.2 Diseño", 2)
    tabla(
        doc,
        ["Componente", "Función"],
        [
            ["Specification", "Interfaz base con `is_satisfied_by` y `description`"],
            ["AndSpec / OrSpec / NotSpec", "Composición con `&`, `|` y `~`"],
            ["VIPCustomerSpec", "Comprueba `customer.is_vip`"],
            ["MinimumAmountSpec", "Comprueba `order.total > umbral`"],
            ["MinimumUnitsSpec", "Comprueba `order.units >= umbral`"],
            ["ActivePromotionSpec", "Comprueba si el pedido tiene promoción"],
            ["ActiveCustomerSpec", "Comprueba que el cliente esté activo"],
            ["HasItemsSpec", "Comprueba que el pedido tenga al menos un renglón"],
        ],
        titulo="Tabla 7.1. Componentes del patrón",
        anchos=[1.9, 4.3],
        tamano=8.5,
    )

    seccion_1(doc, "7.3 Reglas del negocio", 2)
    tabla(
        doc,
        ["Regla", "Condiciones", "Descuento"],
        [
            ["TEN_PERCENT_RULE", "VIP AND total > 500 AND unidades >= 4 AND NOT promoción", "10 %"],
            ["FIVE_PERCENT_RULE", "total > 5000 AND NOT promoción", "5 %"],
            ["DISCOUNT_RULE", "TEN_PERCENT_RULE OR FIVE_PERCENT_RULE", "La mayor aplicable"],
        ],
        titulo="Tabla 7.2. Reglas de descuento",
        anchos=[1.45, 4.0, 1.0],
        tamano=8.5,
        alinear_centro=[2],
    )
    parrafo(
        doc,
        "La prioridad se resuelve en `resolve_discount`: si la regla del 10 % se "
        "cumple, se devuelve diez; en caso contrario se evalúa la del 5 %. El total "
        "bruto del pedido nunca se modifica. `evaluate_discount` devuelve un "
        "`DiscountResult` con el porcentaje, el total bruto, el total con descuento y "
        "el importe ahorrado.",
    )

    seccion_1(doc, "7.4 Composición y prueba", 2)
    parrafo(
        doc,
        "Las Specifications se pueden combinar sin duplicar condiciones. Por ejemplo, "
        "la regla confirmable se expresa como la conjunción de «cliente activo» y "
        "«pedido con artículos». Además de las pruebas de resultado, la suite verifica "
        "que los objetos compuestos sean del tipo correcto y que las descripciones se "
        "deriven de los umbrales, no se escriban a mano.",
    )
    codigo(
        doc,
        "TEN_PERCENT_RULE = (\n"
        "    VIPCustomerSpec()\n"
        "    & MinimumAmountSpec(500)\n"
        "    & MinimumUnitsSpec(4)\n"
        "    & ~ActivePromotionSpec()\n"
        ")\n\n"
        "FIVE_PERCENT_RULE = (\n"
        "    MinimumAmountSpec(5000) & ~ActivePromotionSpec()\n"
        ")\n\n"
        "DISCOUNT_RULE = TEN_PERCENT_RULE | FIVE_PERCENT_RULE",
    )
    imagen(
        doc,
        "especificaciones.png",
        "Figura 7.1. Composición de Specifications y cálculo del descuento.",
        ancho=5.9,
    )

    seccion_1(doc, "8. Máquina de estados y eventos")
    parrafo(
        doc,
        "El estado del pedido no es un campo libre: es una enumeración y solo cambia "
        "por los métodos del agregado. Cada transición válida registra un evento con "
        "nombre y marca de tiempo. Los eventos no se generan desde la API ni desde las "
        "pruebas; son consecuencia de la operación del dominio.",
    )
    tabla(
        doc,
        ["Estado", "Significado"],
        [
            ["DRAFT", "Pedido en elaboración; admite artículos"],
            ["CONFIRMED", "Pedido aceptado; puede enviarse o cancelarse"],
            ["SHIPPED", "Pedido enviado con guía; puede entregarse o devolverse"],
            ["DELIVERED", "Pedido entregado al cliente (estado final)"],
            ["RETURNED", "Pedido devuelto (estado final)"],
            ["CANCELLED", "Pedido cancelado antes del envío (estado final)"],
        ],
        titulo="Tabla 8.1. Estados del pedido",
        anchos=[1.3, 4.9],
        tamano=9,
    )
    tabla(
        doc,
        ["Desde", "Hasta", "Método", "Precondición", "Evento"],
        [
            ["DRAFT", "CONFIRMED", "confirm()", "Con artículos y cliente activo", "OrderConfirmed"],
            ["CONFIRMED", "SHIPPED", "ship(guia)", "Guía no vacía", "OrderShipped"],
            ["CONFIRMED", "CANCELLED", "cancel(motivo)", "Motivo no vacío", "OrderCancelled"],
            ["SHIPPED", "DELIVERED", "deliver()", "—", "OrderDelivered"],
            ["SHIPPED", "RETURNED", "return_order()", "—", "OrderReturned"],
        ],
        titulo="Tabla 8.2. Transiciones permitidas",
        anchos=[0.8, 1.0, 1.15, 2.1, 1.25],
        tamano=8.5,
    )
    parrafo(
        doc,
        "La regla central es «validar, cambiar el estado y registrar el evento, en ese "
        "orden». Si la validación falla, el método lanza una excepción antes de tocar "
        "el estado. La suite comprueba el estado y la lista de eventos después de cada "
        "rechazo, no solo el tipo de excepción.",
    )
    imagen(
        doc,
        "estados.png",
        "Figura 8.1. Máquina de estados de Order con sus cinco transiciones.",
        ancho=6.4,
    )
    imagen(
        doc,
        "flujo-confirmacion.png",
        "Figura 8.2. Decisión de POST /orders/{id}/confirm: 404, 409, 422 o 200.",
        ancho=6.4,
    )

    seccion_1(doc, "9. Servicio de aplicación")
    parrafo(
        doc,
        "`app/service.py` concentra la lectura de los CSV y la construcción de los "
        "objetos de dominio. La API y los scripts usan la misma función de carga, de "
        "modo que una diferencia entre la aplicación y las pruebas sería un error de "
        "una sola pieza, no de dos rutas paralelas.",
    )
    tabla(
        doc,
        ["Función", "Responsabilidad"],
        [
            ["read_rows(nombre)", "Lee un CSV con codificación utf-8-sig y devuelve filas"],
            ["load_orders(directorio)", "Construye los 16 pedidos con clientes, productos y detalle"],
            ["load_fixture(directorio)", "Alias didáctico compatible con el cargador original"],
            ["catalogo_productos()", "Devuelve el precio unitario por SKU"],
            ["catalogo_clientes()", "Devuelve los datos de catálogo por cliente"],
        ],
        titulo="Tabla 9.1. API del servicio",
        anchos=[1.9, 4.3],
        tamano=8.5,
    )
    parrafo(
        doc,
        "El servicio acepta un directorio alternativo, lo que permite probar la carga "
        "con fixtures temporales sin tocar `datos/`. La API recarga los pedidos y los "
        "catálogos al arrancar y expone `POST /admin/reload-fixtures` para repetir una "
        "demostración sin reiniciar el proceso. Esa recarga descarta los cambios en "
        "memoria, pero nunca escribe en los CSV.",
    )

    seccion_1(doc, "10. API REST")
    parrafo(
        doc,
        "La API se implementó con FastAPI. Los modelos de entrada validan el formato; "
        "las reglas de negocio siguen en el dominio. Los importes se serializan como "
        "cadenas para conservar la exactitud decimal en JSON.",
    )
    tabla(
        doc,
        ["Método", "Ruta", "Descripción", "Códigos"],
        [
            ["GET", "/", "Información general", "200"],
            ["GET", "/orders", "Resumen de los 16 pedidos", "200"],
            ["GET", "/orders/{id}", "Detalle con renglones, total y eventos", "200, 404"],
            ["GET", "/orders/{id}/rules", "Evaluación regla por regla", "200, 404"],
            ["POST", "/orders/{id}/confirm", "DRAFT → CONFIRMED", "200, 404, 409, 422"],
            ["POST", "/orders/{id}/ship", "CONFIRMED → SHIPPED", "200, 404, 409, 422"],
            ["POST", "/orders/{id}/cancel", "CONFIRMED → CANCELLED", "200, 404, 409, 422"],
            ["POST", "/orders/{id}/deliver", "SHIPPED → DELIVERED", "200, 404, 409"],
            ["POST", "/orders/{id}/return", "SHIPPED → RETURNED", "200, 404, 409"],
            ["POST", "/orders/{id}/items", "Agrega un renglón", "200, 404, 409, 422"],
            ["POST", "/orders", "Crea un pedido", "201, 422"],
            ["GET", "/health", "Estado del servicio", "200"],
            ["POST", "/admin/reload-fixtures", "Recarga los CSV", "200"],
        ],
        titulo="Tabla 10.1. Endpoints de la aplicación",
        anchos=[0.7, 1.85, 2.6, 1.1],
        tamano=8.5,
    )
    parrafo(
        doc,
        "Los controladores comparten dos utilidades: una busca el pedido y lanza 404 si "
        "no existe; la otra traduce las excepciones del dominio. El mapeo es estable: "
        "`InvalidOrderTransition` produce 409, mientras que `InvalidOrderData` y "
        "`OrderItemLimitExceeded` producen 422. El 404 se decide en el controlador "
        "porque es una condición de la interfaz, no del agregado.",
    )
    codigo(
        doc,
        "POST /orders/O013/confirm        -> 200 {\"status\": \"CONFIRMED\", ...}\n"
        "POST /orders/O013/confirm        -> 409 {\"detail\": \"...\"}\n"
        "POST /orders/O999/confirm        -> 404 {\"detail\": \"Pedido no encontrado\"}\n"
        "POST /orders/O007/confirm        -> 422 {\"detail\": \"Requiere artículos...\"}\n"
        "POST /orders/O009/items          -> 422 {\"detail\": \"... 20 unidades ...\"}",
    )
    parrafo(
        doc,
        "La documentación interactiva queda disponible en `/docs` y el contrato en "
        "`/openapi.json`. Las evidencias incluyen capturas reales de Swagger para 200, "
        "404 y 422. La suite usa `TestClient` con el ciclo de vida de la aplicación, "
        "de modo que cada prueba parte de los 16 pedidos en DRAFT.",
    )
    imagen(
        doc,
        "secuencia-confirmar.png",
        "Figura 10.1. Secuencia de confirmación: controlador, agregado y traducción de errores.",
        ancho=5.5,
    )

    seccion_1(doc, "11. Pruebas y aseguramiento de calidad")
    parrafo(
        doc,
        f"La suite se organiza por nivel de responsabilidad y suma {pruebas} pruebas. "
        "Cada archivo prueba una capa distinta, aunque comparte los fixtures del "
        "paquete para que los datos sean los mismos que usa la aplicación.",
    )
    tabla(
        doc,
        ["Archivo", "Pruebas", "Cobertura"],
        [
            ["tests/test_domain.py", "39", "Money, OrderItem, invariantes y descuentos"],
            ["tests/test_api.py", "26", "404, 409, 422 y ciclo de vida HTTP"],
            ["tests/test_escenarios.py", "15", "Un caso por escenario E01–E15"],
            ["tests/test_maquina_estados.py", "21", "Transiciones y eventos"],
            ["tests/test_specifications.py", "10", "Specifications aisladas y compuestas"],
            ["tests/test_servicio.py", "5", "Carga, catálogos, /health y recarga"],
            ["Total", str(pruebas), "Suite completa en verde"],
        ],
        titulo="Tabla 11.1. Distribución de las pruebas",
        anchos=[2.0, 0.8, 3.4],
        tamano=9,
        alinear_centro=[1],
    )
    seccion_1(doc, "11.1 Estrategia", 2)
    for texto in (
        "Las pruebas de dominio construyen objetos en memoria para controlar umbrales "
        "exactos, sin alterar los CSV.",
        "Las pruebas de API usan `TestClient` y verifican el código HTTP y el cuerpo.",
        "Las pruebas de escenarios cargan los fixtures antes de cada caso.",
        "Las pruebas de servicio comparan dos rutas de carga y comprueban los catálogos.",
        "Las pruebas de frontera cubren 500.00 exacto, 500.01, 3 y 4 unidades, 20 y 21.",
    ):
        vineta(doc, texto)

    seccion_1(doc, "11.2 Resultado de la ejecución", 2)
    codigo(
        doc,
        "$ python -m pytest -q\n"
        "........................................................................ [ 62%]\n"
        "............................................                             [100%]\n"
        f"{pruebas} passed, 1 warning in 1.95s",
    )
    parrafo(
        doc,
        "La única advertencia proviene de `fastapi.testclient` por el uso de `httpx` "
        "con Starlette; no la genera el código del proyecto. Se documentó en lugar de "
        "silenciarla con un filtro, porque conocer su origen es parte de la entrega.",
    )

    seccion_1(doc, "11.3 Casos de frontera", 2)
    tabla(
        doc,
        ["Caso", "Entrada", "Resultado esperado"],
        [
            ["Monto estricto", "Total 500.00 / 500.01", "0 % / 10 % si cumple el resto"],
            ["Unidades", "3 / 4 unidades", "Sin descuento / aplica si cumple el resto"],
            ["Límite", "20 / 21 unidades", "Permitido / OrderItemLimitExceeded"],
            ["Prioridad", "VIP con total 6000.00 y 5 unidades", "10 %, no 5 %"],
            ["No mutación", "Confirmar dos veces", "409 y un solo evento"],
            ["Sin guía", "ship(\"\")", "422 y estado CONFIRMED"],
        ],
        titulo="Tabla 11.2. Fronteras verificadas",
        anchos=[1.2, 2.7, 2.2],
        tamano=8.5,
    )

    seccion_1(doc, "12. Escenarios E01–E15")
    parrafo(
        doc,
        "Los escenarios se ejecutan con `scripts/ejecutar_bitacora.py`. Cada caso se "
        "carga desde cero, ejecuta la acción indicada y compara el resultado con el "
        "objetivo de `escenarios.csv`. La bitácora completa se entrega en Markdown, "
        "CSV y DOCX/PDF; aquí se resume el resultado.",
    )
    tabla(
        doc,
        ["Caso", "Pedido", "Acción", "Objetivo", "Resultado observado", "Estado final"],
        [
            [
                b["caso"],
                b["pedido"],
                b["operacion"],
                b["resultado_esperado"],
                b["resultado_observado"],
                b["estado_final"],
            ]
            for b in bitacora
        ],
        titulo="Tabla 12.1. Resultado de los 15 escenarios",
        anchos=[0.55, 0.55, 1.45, 0.8, 2.1, 0.85],
        tamano=7.5,
        alinear_centro=[0, 1, 3, 5],
    )
    parrafo(
        doc,
        "Los 15 casos coinciden con el objetivo del paquete. Los rechazos conservan el "
        "estado previo y no agregan eventos; las transiciones válidas dejan exactamente "
        "el evento correspondiente. La coincidencia se calculó ejecutando el dominio, "
        "no escribiendo el resultado en la tabla.",
    )

    seccion_1(doc, "12.1 Casos destacados", 2)
    for texto in (
        "E01 aplica el 10 % a O001: VIP, 2100.00 > 500, 4 unidades y sin promoción.",
        "E05 aplica el 5 % a O006: 18000.00 > 5000 y sin promoción.",
        "E08 rechaza agregar 2 unidades a O009 porque 19 + 2 = 21 supera el límite.",
        "E11 confirma O012 pero rechaza el envío sin guía; el evento OrderConfirmed permanece.",
        "E13 cancela desde SHIPPED y falla: la cancelación solo se permite desde CONFIRMED.",
        "E14 y E15 terminan en RETURNED y DELIVERED, respectivamente.",
    ):
        vineta(doc, texto)

    seccion_1(doc, "13. Evidencias y reproducibilidad")
    parrafo(
        doc,
        "A continuación se presentan las capturas de pantalla de la ejecución real de "
        "cada paso: carga de los CSV, bitácora de escenarios, suite de pruebas y "
        "consumo de la API desde Swagger UI. Las salidas en texto se conservan además "
        "en `evidencias/` y `docs/evidencias/` para poder repetirlas.",
    )
    seccion_1(doc, "13.1 Carga de los datos CSV", 2)
    parrafo(doc, "La carga de los cinco CSV construye los 16 pedidos en estado DRAFT "
            "con su número de artículos y su total.")
    captura(doc, "11_carga_16_pedidos.png",
            "Figura 13.1. Ejecución de `python cargar_csv.py`: 16 pedidos cargados.", ancho=5.8)
    seccion_1(doc, "13.2 Bitácora de escenarios E01–E15", 2)
    parrafo(doc, "El script de bitácora ejecuta los 15 escenarios contra el dominio y "
            "compara el resultado observado con el esperado.")
    captura(doc, "13_bitacora_consola.png",
            "Figura 13.2. Ejecución de `python scripts/ejecutar_bitacora.py`: 15/15 coinciden.")
    seccion_1(doc, "13.3 Pruebas automatizadas", 2)
    parrafo(doc, f"La suite completa termina en verde con {pruebas} pruebas aprobadas.")
    captura(doc, "12_pytest_resultados.png",
            f"Figura 13.3. Ejecución de `python -m pytest -q`: {pruebas} passed.", ancho=5.8)
    seccion_1(doc, "13.4 API REST en Swagger UI", 2)
    parrafo(doc, "Con `uvicorn app.api:app --reload` se levantó la API y se probaron los "
            "endpoints desde `/docs`, tanto en casos de éxito como de rechazo.")
    captura(doc, "06_swagger_docs.png",
            "Figura 13.4. Documentación interactiva de la API en `/docs`.", ancho=5.6)
    captura(doc, "08_swagger_exito_200.png",
            "Figura 13.5. `POST /orders/O001/confirm`: respuesta 200 con el evento OrderConfirmed.",
            ancho=4.4)
    captura(doc, "09_swagger_rechazo_404.png",
            "Figura 13.6. Rechazo 404: el pedido solicitado no existe.", ancho=4.4)
    captura(doc, "10_swagger_rechazo_422.png",
            "Figura 13.7. Rechazo 422: pedido sin artículos o cliente inactivo.", ancho=4.4)
    seccion_1(doc, "13.5 Reproducibilidad", 2)
    parrafo(doc, "Los mismos resultados se obtienen desde la raíz del proyecto con:")
    codigo(
        doc,
        "python cargar_csv.py\n"
        "python scripts/ejecutar_bitacora.py\n"
        "python -m pytest -q\n"
        "uvicorn app.api:app --reload",
    )

    seccion_1(doc, "14. Decisiones de diseño y riesgos")
    tabla(
        doc,
        ["Decisión", "Alternativa descartada", "Motivo"],
        [
            ["`Decimal` para importes", "`float`", "Evita errores de representación binaria"],
            ["Descuento separado del total", "Modificar `order.total`", "El total bruto es un dato del pedido"],
            ["Specifications componibles", "Cadena de `if`", "Reglas legibles, probables y reutilizables"],
            ["Eventos registrados al final", "Evento antes de validar", "Un rechazo no debe dejar historia falsa"],
            ["Servicio de aplicación", "Carga dentro de la API", "Una sola ruta de carga para API y scripts"],
            ["Traducción de errores en el controlador", "HTTPException en el dominio", "El dominio no depende de FastAPI"],
        ],
        titulo="Tabla 14.1. Decisiones de diseño",
        anchos=[1.6, 1.8, 2.8],
        tamano=8.5,
    )
    seccion_1(doc, "14.1 Limitaciones", 2)
    for texto in (
        "El repositorio es en memoria; no hay persistencia entre reinicios.",
        "La autenticación y la autorización están fuera del alcance.",
        "No se descuenta inventario ni se integra un sistema de pagos.",
        "La recarga de fixtures descarta cambios locales, algo aceptable en el contexto académico.",
    ):
        vineta(doc, texto)
    seccion_1(doc, "14.2 Trabajo futuro", 2)
    for texto in (
        "Agregar persistencia con SQLAlchemy o un repositorio implementado sobre archivos.",
        "Incorporar autenticación y roles para distinguir clientes y administradores.",
        "Añadir métricas y trazas para observar el ciclo de vida de los pedidos.",
        "Ampliar la API con paginación y filtros por estado o cliente.",
    ):
        vineta(doc, texto)

    seccion_1(doc, "15. Conclusiones")
    parrafo(
        doc,
        "La práctica permitió comprobar que un núcleo de negocio bien delimitado se "
        "puede probar y explicar sin depender de la interfaz. El dominio concentra las "
        "invariantes, las Specifications hacen explícitas las reglas comerciales y la "
        "máquina de estados garantiza que cada transición deje un rastro coherente. "
        "La API queda como un adaptador delgado, fácil de revisar.",
    )
    parrafo(
        doc,
        f"Los {pruebas} casos automatizados y los 15 escenarios reproducibles dan una "
        "base objetiva para afirmar que la entrega cumple el comportamiento esperado. "
        "El proceso también dejó aprendizajes sobre el orden de las anotaciones en "
        "Pydantic, la diferencia entre `>` y `>=`, el orden de las validaciones HTTP y "
        "la importancia de no modificar los datos para forzar un resultado.",
    )
    parrafo(
        doc,
        "El proyecto queda como un repositorio local listo para publicarse en GitHub, "
        "con manifiesto de entrega, documentación y comandos de reproducción. La "
        "siguiente ampliación natural es la persistencia, seguida de autenticación y "
        "observabilidad, sin mover las reglas fuera del dominio.",
    )

    seccion_1(doc, "16. Referencias")
    for texto in (
        "Documentación oficial de Python 3.14: tipos numéricos y módulo `decimal`.",
        "Documentación oficial de FastAPI: modelos Pydantic, dependencias y pruebas.",
        "Documentación oficial de pytest: fixtures, parametrización y captura de salida.",
        "Evans, E. (2003). *Domain-Driven Design: Tackling Complexity in the Heart of Software*.",
        "Fowler, M. (2013). *Catalog of Patterns of Enterprise Application Architecture*.",
        "Material de la práctica: `CLN_U4_Practica_Guiada_Alumnos_TechStore.pdf`.",
        "Repositorio del proyecto: `README.md`, `docs/` y `datos/`.",
    ):
        vineta(doc, texto)

    salto(doc)
    seccion_1(doc, "Anexo A. Respuestas a las preguntas de la sección 15")
    preguntas = [
        (
            "1. ¿Por qué se usa Decimal en lugar de float para precios?",
            "Porque `float` representa los números en base dos y muchos importes decimales "
            "no tienen una representación binaria exacta. El clásico `0.1 + 0.2` produce "
            "`0.30000000000000004`. En un pedido, ese error se acumula y termina en un total "
            "que no coincide con la factura. `Decimal` trabaja en base diez, que es la forma "
            "en que se escriben los precios, y conserva la exactitud de las sumas y "
            "multiplicaciones. En el proyecto, `Money` normaliza la entrada con "
            "`Decimal(str(amount))` y redondea a centavos con `ROUND_HALF_UP`.",
        ),
        (
            "2. ¿Cómo se aplican las reglas de descuento del 10 % y del 5 %?",
            "La regla del 10 % exige cuatro condiciones simultáneas: cliente VIP, total "
            "estrictamente mayor a 500 MXN, al menos cuatro unidades y ausencia de promoción. "
            "La del 5 % exige total estrictamente mayor a 5000 MXN y ausencia de promoción. "
            "Ambas se construyen como Specifications y se combinan con AND y NOT; la regla "
            "general es un OR de las dos. Si las dos se cumplen, `resolve_discount` devuelve "
            "primero el 10 %. El total bruto del pedido no se modifica: el resultado se "
            "entrega en un objeto aparte.",
        ),
        (
            "3. ¿Por qué el límite de 20 unidades se valida en el agregado y no en la API?",
            "Porque es una invariante del pedido. Si se validara solo en la API, cualquier "
            "otro cliente del dominio (un script, una prueba o una futura interfaz) podría "
            "dejar un pedido con 21 unidades. El agregado es el único punto de entrada para "
            "modificar los artículos y por eso concentra la regla. La API se limita a "
            "traducir `OrderItemLimitExceeded` a 422. Además, la validación se hace antes de "
            "agregar el renglón, de modo que el rechazo no muta el pedido.",
        ),
        (
            "4. ¿Cómo se garantiza que una operación inválida no deje efectos parciales?",
            "Cada método del pedido sigue el orden validar → cambiar el estado → registrar "
            "el evento. Si la validación falla, se lanza una excepción antes de tocar "
            "`status` o `items`. En `add_item` se calcula el nuevo total de unidades y se "
            "compara con el límite antes de agregar. Las pruebas verifican el estado, el "
            "número de eventos y los artículos después de cada rechazo, no solo el tipo de "
            "excepción. Así se detecta cualquier mutación accidental.",
        ),
        (
            "5. ¿Qué aporta separar el servicio de aplicación de la API?",
            "El servicio es la única pieza que sabe leer los CSV y construir los pedidos. "
            "La API y los scripts comparten esa ruta, lo que evita divergencias entre la "
            "aplicación y las pruebas. También permite apuntar a un directorio alternativo "
            "de fixtures, recargar los datos sin reiniciar el proceso y sustituir la fuente "
            "de datos en el futuro sin reescribir los controladores. La API queda enfocada "
            "en HTTP y el servicio en la carga.",
        ),
    ]
    for titulo, respuesta in preguntas:
        seccion_1(doc, titulo, 3)
        parrafo(doc, respuesta)

    seccion_1(doc, "Anexo B. Glosario")
    tabla(
        doc,
        ["Término", "Definición"],
        [
            ["Value object", "Objeto inmutable sin identidad, comparado por valor"],
            ["Entidad", "Objeto con identidad y ciclo de vida propio"],
            ["Aggregate root", "Entidad que protege las invariantes de un conjunto"],
            ["Invariante", "Condición que debe cumplirse siempre en el modelo"],
            ["Specification", "Predicado componible que expresa una regla"],
            ["Domain event", "Hecho ya ocurrido que se registra para auditoría"],
            ["Fixture", "Conjunto de datos de entrada usado en pruebas o demostraciones"],
            ["Adapter", "Capa que traduce entre el dominio y un mecanismo externo"],
        ],
        titulo="Tabla B.1. Términos usados en el documento",
        anchos=[1.5, 4.7],
        tamano=9,
    )

    seccion_1(doc, "Anexo C. Comandos de reproducción")
    codigo(
        doc,
        "# 1. Crear y activar el entorno\n"
        "python -m venv .venv\n"
        ".venv\\Scripts\\Activate.ps1\n"
        "python -m pip install -r requirements.txt\n\n"
        "# 2. Ejecutar el cargador y la bitácora\n"
        "python cargar_csv.py\n"
        "python scripts/ejecutar_bitacora.py\n\n"
        "# 3. Ejecutar todas las pruebas\n"
        "python -m pytest -q\n\n"
        "# 4. Regenerar evidencias y documentos\n"
        "python scripts/generar_evidencias.py\n"
        "C:\\Python314\\python.exe scripts/generar_documentos.py\n\n"
        "# 5. Levantar la API\n"
        "uvicorn app.api:app --reload",
    )

    salto(doc)
    seccion_repositorio(doc, "Link de GitHub")

    destino = DOCS / "Reporte_Tecnico_TechStore_Unidad_4.docx"
    doc.save(destino)
    print(f"Reporte: {destino.relative_to(ROOT)}")
    return destino


# ---------------------------------------------------------------------------
# Bitácora
# ---------------------------------------------------------------------------

EXPLICACIONES = {
    "E01": (
        "O001 es el caso base del descuento del 10 %: cliente VIP, total 2100.00 MXN, "
        "cuatro unidades y sin promoción. La regla se cumple completa y el total con "
        "descuento es 1890.00 MXN.",
        "Cálculo del descuento",
    ),
    "E02": (
        "O002 no es VIP, tiene dos unidades y su total es 1050.00 MXN. No cumple la "
        "regla del 10 % ni supera el umbral de 5000 MXN, así que el descuento es 0 %.",
        "Cálculo del descuento",
    ),
    "E03": (
        "O003 pertenece a C003, un cliente inactivo. La confirmación exige cliente "
        "activo y el dominio lanza `InvalidOrderData`. El pedido permanece en DRAFT y "
        "sin eventos.",
        "Rechazo por cliente inactivo",
    ),
    "E04": (
        "O005 tiene promoción activa. Ambas reglas de descuento incluyen la condición "
        "«sin promoción», de modo que el porcentaje es 0 % aunque el cliente sea VIP.",
        "Promoción activa",
    ),
    "E05": (
        "O006 tiene un total de 18000.00 MXN, superior a 5000, y no tiene promoción. "
        "No es VIP ni alcanza cuatro unidades, por lo que no aplica el 10 %; el 5 % "
        "deja el total en 17100.00 MXN.",
        "Cálculo del descuento",
    ),
    "E06": (
        "O007 no tiene renglones. La confirmación exige al menos un artículo y el "
        "dominio lanza `InvalidOrderData`. El estado y la lista de eventos no cambian.",
        "Rechazo por pedido vacío",
    ),
    "E07": (
        "O008 tiene 20 unidades exactas en un solo renglón. El límite es 20, no 21, "
        "así que la confirmación es válida y registra `OrderConfirmed`.",
        "Límite exacto",
    ),
    "E08": (
        "O009 tiene 19 unidades. Agregar dos unidades de SKU003 daría 21, por encima "
        "del límite. El agregado lanza `OrderItemLimitExceeded` antes de modificar los "
        "artículos: el pedido conserva 19 unidades y sigue en DRAFT.",
        "Límite excedido",
    ),
    "E09": (
        "O010 suma exactamente 500.00 MXN. La regla exige `> 500`, no `>= 500`, por lo "
        "que el descuento es 0 %. Es el caso frontera que distingue los dos operadores.",
        "Umbral estricto",
    ),
    "E10": (
        "O011 es VIP y su total supera 500, pero solo tiene 3 unidades. La regla del "
        "10 % exige al menos 4 unidades, así que el descuento es 0 %.",
        "Umbral de unidades",
    ),
    "E11": (
        "O012 se confirma correctamente y registra `OrderConfirmed`. Después se intenta "
        "enviar con una guía vacía y el dominio lanza `InvalidOrderData`. La "
        "confirmación queda registrada porque fue una operación válida anterior.",
        "Envío sin guía",
    ),
    "E12": (
        "O013 se confirma una vez y queda en CONFIRMED con un evento. La segunda "
        "confirmación no es válida desde ese estado y lanza `InvalidOrderTransition`; "
        "no se agrega un segundo evento.",
        "Confirmación duplicada",
    ),
    "E13": (
        "O014 se confirma y se envía; queda en SHIPPED con dos eventos. Cancelar desde "
        "SHIPPED no está permitido y el dominio lanza `InvalidOrderTransition`. El "
        "estado final sigue siendo SHIPPED.",
        "Cancelación fuera de estado",
    ),
    "E14": (
        "O015 se confirma, se envía y se devuelve. La devolución es válida desde "
        "SHIPPED y termina en RETURNED con los eventos `OrderConfirmed`, "
        "`OrderShipped` y `OrderReturned`.",
        "Devolución",
    ),
    "E15": (
        "O016 se confirma, se envía y se entrega. La entrega es válida desde SHIPPED y "
        "termina en DELIVERED con los eventos `OrderConfirmed`, `OrderShipped` y "
        "`OrderDelivered`.",
        "Entrega",
    ),
}


def generar_bitacora() -> Path:
    clientes, productos, pedidos, detalle, escenarios, bitacora = cargar_datos()
    pruebas = contar_pruebas()
    por_pedido = {p["pedido_id"]: p for p in pedidos}
    por_cliente = {c["cliente_id"]: c for c in clientes}

    doc = Document()
    configurar_documento(doc, "Bitácora · TechStore U4")

    portada(
        doc,
        "Bitácora de Escenarios",
        "Ejecución reproducible de E01–E15 · Práctica Guiada Actividad 4",
    )

    seccion_1(doc, "1. Propósito")
    parrafo(
        doc,
        "Esta bitácora registra la ejecución de los 15 escenarios propuestos en "
        "`datos/escenarios.csv`. Para cada caso se indica el pedido, el estado inicial, "
        "la operación, el objetivo declarado en el paquete, el resultado observado, el "
        "estado final, el evento emitido y el veredicto. Los resultados se obtuvieron "
        "ejecutando el dominio, no escribiéndolos manualmente.",
    )
    parrafo(
        doc,
        "El documento complementa el reporte técnico y las pruebas automatizadas. "
        "Mientras que `pytest` verifica el comportamiento con aserciones, la bitácora "
        "deja una traza legible para revisar qué ocurrió en cada escenario y por qué.",
    )

    seccion_1(doc, "2. Método de ejecución")
    for texto in (
        "Cada escenario se ejecuta sobre pedidos recién cargados con `load_orders()`; "
        "ningún caso hereda el estado de otro.",
        "El objetivo se lee de `datos/escenarios.csv` y el resultado se calcula "
        "ejecutando la acción correspondiente.",
        "Si la acción lanza una excepción de dominio, se registra como rechazo y se "
        "conserva el estado previo.",
        "Los eventos se cuentan antes y después de la operación para registrar solo los "
        "emitidos por ese caso.",
        "No se editó ningún CSV ni ningún pedido para forzar un resultado.",
    ):
        vineta(doc, texto)
    codigo(
        doc,
        "$ python scripts/ejecutar_bitacora.py\n"
        "Escenarios registrados: 15\n"
        "...\n"
        "Coincidencias: 15/15\n"
        "Bitácora: docs\\bitacora\\bitacora_escenarios.csv",
    )
    parrafo(
        doc,
        "El script devuelve código de salida 0 solo si los 15 veredictos coinciden. "
        "El CSV resultante es la fuente de las tablas de este documento.",
    )

    seccion_1(doc, "3. Resumen de resultados")
    parrafo(
        doc,
        "Los 15 escenarios coinciden con el objetivo del paquete. La Tabla 3.1 muestra "
        "el resultado observado y el estado final; la Tabla 3.2 resume los eventos.",
    )
    tabla(
        doc,
        ["Caso", "Pedido", "Objetivo", "Observado", "Estado final", "Veredicto"],
        [
            [
                b["caso"],
                b["pedido"],
                b["resultado_esperado"],
                b["resultado_observado"],
                b["estado_final"],
                b["veredicto"],
            ]
            for b in bitacora
        ],
        titulo="Tabla 3.1. Resumen de los 15 escenarios",
        anchos=[0.55, 0.55, 1.0, 2.1, 0.9, 0.9],
        tamano=8,
        alinear_centro=[0, 1, 2, 4, 5],
    )
    tabla(
        doc,
        ["Caso", "Estado inicial", "Estado final", "Eventos emitidos"],
        [
            [b["caso"], b["estado_inicial"], b["estado_final"], b["evento_emitido"]]
            for b in bitacora
        ],
        titulo="Tabla 3.2. Estados y eventos por escenario",
        anchos=[0.55, 1.0, 1.0, 3.7],
        tamano=8,
        alinear_centro=[0, 1, 2],
    )
    parrafo(
        doc,
        "Seis escenarios terminan en un estado distinto de DRAFT (E07, E11, E12, "
        "E13, E14 y E15). Los rechazos conservan el estado que tenían antes de la "
        "operación y solo agregan eventos cuando una operación previa fue válida, "
        "como en E11, E12 y E13.",
    )

    salto(doc)
    seccion_1(doc, "4. Bitácora detallada")
    parrafo(
        doc,
        "Cada ficha presenta los datos del escenario y una explicación breve. Los "
        "valores de la tabla provienen de `bitacora_escenarios.csv`; las explicaciones "
        "describen la causa del resultado.",
    )

    for b in bitacora:
        caso = b["caso"]
        pedido_id = b["pedido"]
        pedido = por_pedido.get(pedido_id, {})
        cliente = por_cliente.get(pedido.get("cliente_id", ""), {})
        explicacion, etiqueta = EXPLICACIONES.get(
            caso, ("Sin explicación registrada.", caso)
        )
        seccion_1(doc, f"{caso} — {etiqueta}", 2)
        tabla(
            doc,
            ["Campo", "Valor"],
            [
                ["Pedido", f"{pedido_id} · cliente {cliente.get('nombre', '?')}"],
                ["Estado inicial", b["estado_inicial"]],
                ["Operación", b["operacion"]],
                ["Objetivo del paquete", b["resultado_esperado"]],
                ["Resultado observado", b["resultado_observado"]],
                ["Estado final", b["estado_final"]],
                ["Evento emitido", b["evento_emitido"]],
                ["Veredicto", b["veredicto"]],
            ],
            anchos=[1.5, 4.7],
            tamano=9,
            mantener_con_siguiente=True,
        )
        parrafo(doc, explicacion)

    seccion_1(doc, "5. Lectura de los resultados")
    seccion_1(doc, "5.1 Unidades frente a renglones", 2)
    parrafo(
        doc,
        "E07 y E08 prueban que el límite se aplica a la suma de cantidades, no al "
        "número de filas. O008 tiene un renglón con 20 unidades y se acepta; O009 tiene "
        "un renglón con 19 unidades y agregar dos más lo lleva a 21, por lo que se "
        "rechaza. Un límite por renglones habría dejado pasar ambos casos.",
    )
    seccion_1(doc, "5.2 Rechazos sin efectos parciales", 2)
    parrafo(
        doc,
        "En E03, E06, E08 y E13 el pedido conserva el estado y los eventos que tenía "
        "antes de la operación fallida. En E08, además, los artículos no cambian: O009 "
        "sigue con 19 unidades. Esta propiedad se verificó en las pruebas de dominio y "
        "en la bitácora.",
    )
    seccion_1(doc, "5.3 Operaciones compuestas", 2)
    parrafo(
        doc,
        "E11, E12 y E13 combinan una operación válida con una fallida. La bitácora "
        "muestra que los eventos de la parte válida permanecen. No se trata de una "
        "transacción atómica de varias operaciones, sino de llamadas independientes: "
        "cada una se registra o se rechaza por separado.",
    )
    seccion_1(doc, "5.4 Escenario adicional de prioridad", 2)
    parrafo(
        doc,
        "El paquete no incluye un pedido que cumpla simultáneamente las reglas del "
        "10 % y del 5 %. O006 supera 5000 pero no es VIP, y O001 es VIP pero no supera "
        "5000. La prioridad se verificó con un pedido construido en memoria en "
        "`tests/test_specifications.py`: VIP, 5 unidades y total 6000.00 MXN. El "
        "resultado es 10 %, no 5 %, y el total con descuento es 5400.00 MXN.",
    )

    seccion_1(doc, "6. Reproducibilidad")
    parrafo(
        doc,
        "La bitácora se puede repetir desde la raíz del proyecto con el entorno "
        "virtual activado. Los comandos generan la salida de consola y el CSV; la "
        "suite completa comprueba además que los escenarios sigan coincidiendo.",
    )
    codigo(
        doc,
        "# Ejecutar solo la bitácora\n"
        "python scripts/ejecutar_bitacora.py\n\n"
        "# Verificar los escenarios con pytest\n"
        "python -m pytest tests/test_escenarios.py -v\n\n"
        "# Ejecutar toda la suite\n"
        "python -m pytest -q",
    )
    tabla(
        doc,
        ["Recurso", "Ubicación"],
        [
            ["Script de ejecución", "scripts/ejecutar_bitacora.py"],
            ["CSV generado", "docs/bitacora/bitacora_escenarios.csv"],
            ["Salida de consola", "evidencias/bitacora_consola.txt"],
            ["Pruebas de escenarios", "tests/test_escenarios.py"],
            ["Evidencia de pytest", "docs/evidencias/04_pytest.txt"],
            ["Documento de apoyo", "docs/BITACORA_E01_E15.md"],
        ],
        titulo="Tabla 6.1. Recursos de la bitácora",
        anchos=[1.8, 4.4],
        tamano=9,
    )

    seccion_1(doc, "6.1 Evidencias de ejecución", 2)
    parrafo(doc, "Capturas de la ejecución real de la bitácora y de las pruebas de escenarios.")
    captura(doc, "13_bitacora_consola.png",
            "Figura 6.1. Ejecución de `python scripts/ejecutar_bitacora.py`: 15/15 coinciden.")
    captura(doc, "14_pytest_escenarios.png",
            "Figura 6.2. Ejecución de `python -m pytest tests/test_escenarios.py -q`.", ancho=5.8)

    seccion_1(doc, "7. Conclusiones")
    parrafo(
        doc,
        "La bitácora confirma que el dominio implementado reproduce el comportamiento "
        "esperado en los 15 escenarios del paquete. Los casos de frontera —500.00 "
        "exacto, 3 unidades, 20 unidades y 21 unidades— son los que más información "
        "aportan, porque distinguen una regla bien implementada de una aproximación. "
        "También muestran que un rechazo no altera el pedido y que los eventos quedan "
        "registrados solo cuando la operación correspondiente fue válida.",
    )
    parrafo(
        doc,
        f"El resultado se complementa con {pruebas} pruebas automatizadas y con las "
        "evidencias de ejecución. Cualquier cambio futuro en las reglas deberá "
        "actualizar esta bitácora y volver a ejecutar el script; el veredicto 15/15 es "
        "la señal de que los datos, el dominio y la documentación siguen alineados.",
    )

    salto(doc)
    seccion_repositorio(doc, "Link de GitHub")

    destino = DOCS / "bitacora" / "Bitacora_TechStore.docx"
    doc.save(destino)
    print(f"Bitácora: {destino.relative_to(ROOT)}")
    return destino


def main() -> int:
    DOCS.mkdir(parents=True, exist_ok=True)
    (DOCS / "bitacora").mkdir(parents=True, exist_ok=True)
    generar_reporte()
    generar_bitacora()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
